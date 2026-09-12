"""Read the authoritative workbook. Never execute workbook instructions/formulas."""
import argparse,hashlib,json,math
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parent

def main():
 p=argparse.ArgumentParser();p.add_argument('workbook',type=Path);a=p.parse_args()
 w=openpyxl.load_workbook(a.workbook,data_only=True)
 def num(v,where):
  if not isinstance(v,(int,float)) or not math.isfinite(v):raise ValueError(f'Missing numeric projection: {where}: {v}')
  return v
 fields=['goals','assists','power_play_points','short_handed_points','shots_on_goal','blocks','hits','penalty_minutes']
 weights={'skater':{},'goalie':{}}
 for key,row in zip(fields,[5,6,7,8,9,10,11,12]):weights['skater'][key]=num(w['Scoring Settings'].cell(row,2).value,f'scoring B{row}')
 for key,row in zip(['wins','shutouts','saves','goals_against'],[15,16,17,18]):weights['goalie'][key]=num(w['Scoring Settings'].cell(row,2).value,f'scoring B{row}')
 players=[]
 for sheet,goalie in [('Draft Board',False),('Goalies',True)]:
  for r,row in enumerate(w[sheet].iter_rows(min_row=2,values_only=True),2):
   if not row[2]:continue
   stats={key:num(row[col-1],f'{sheet}!{r}:{col}') for key,col in (zip(['wins','saves','shutouts','goals_against'],[11,12,13,14]) if goalie else zip(fields,[12,13,16,17,15,19,18,20]))}
   players.append({'key':f'{sheet}:{r}','name':row[2],'team':row[3],'position':'G' if goalie else row[4],'isGoalie':goalie,'source':row[4] if goalie else row[5],'tier':row[5] if goalie else row[6],'baseGames':num(row[6] if goalie else row[7],f'{sheet}!GP{r}'),'games':num(row[7] if goalie else row[8],f'{sheet}!used{r}'),'stats':stats,'rosterProbability':None if goalie else row[23],'line':None if goalie else row[25],'powerPlay':None if goalie else row[26],'note':row[9] if goalie else row[10],'confidence':row[8] if goalie else row[9],'sourceRank':row[1],'sourceFantasyPoints':num(row[17] if goalie else row[21],f'{sheet}!FPTS{r}'),'sourceRow':r})
 assert len({p['name'] for p in players})==len(players),'Ambiguous player names in workbook'
 rookies=[];narrative=[];tier=''
 for r,row in enumerate(w['Rookies'].iter_rows(values_only=True),1):
  if isinstance(row[0],str) and row[0].startswith('TIER '):tier=row[0]
  if row[2] in ['C','LW','RW','D','G'] and row[0]!='Player':
   rookies.append({'name':row[0],'team':row[1],'position':row[2],'nhlGames':row[3],'draft':row[4],'tier':tier,'support':row[9] or '', 'sourceRow':r})
  elif row[0] and row[0]!='Player':
   narrative.append({'row':r,'cells':[str(v) for v in row if v is not None]})
 teams=[]
 for s in w:
  if len(s.title)!=3 or not s.title.isupper():continue
  sections=[];current=None
  for r,row in enumerate(s.iter_rows(values_only=True),1):
   if row[0] and all(v is None for v in row[1:]):
    if r>3:current={'title':str(row[0]),'rows':[]};sections.append(current)
   elif current and row[2] and row[2]!='Player':
    current['rows'].append({'slot':str(row[0] or ''),'position':str(row[1] or ''),'name':str(row[2]),'note':str(row[17] or '') if len(row)>17 else ''})
  teams.append({'team':s.title,'title':s.cell(1,1).value,'intro':s.cell(2,1).value,'sections':sections,'rawRows':[[v for v in row] for row in s.values]})
 data={'source':{'name':a.workbook.name,'sha256':hashlib.sha256(a.workbook.read_bytes()).hexdigest(),'sheets':w.sheetnames},'weights':weights,'seasonGames':w['Scoring Settings']['B21'].value,'players':players,'rookies':rookies,'rookieNarrative':narrative,'teams':teams,'readMe':[r[0] for r in w['Read Me'].values if r[0]],'audit':[r[0] for r in w['Audit'].values if r[0]],'depthCharts':[[v for v in r] for r in w['Depth Charts'].values]}
 (ROOT/'workbook-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
 print(f'Imported {len(players)} players, {len(rookies)} individually profiled rookies, {len(teams)} team tabs. Source {data["source"]["sha256"]}')
if __name__=='__main__':main()
