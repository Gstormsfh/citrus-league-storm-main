"""Canonical workbook export guards and cached formula checks."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl
from export_canonical_xlsx import export, payload
from test_canonical_import import fixtures


class CanonicalWorkbookTests(unittest.TestCase):
    def test_payload_keeps_null_exposure_and_binds_selected_weights(self):
        d, e = fixtures()
        weights = deepcopy(e['weights']);weights['skater']['goals'] = -4
        x = payload(d, e, d['revision'], 'My league', weights)
        self.assertEqual(x['league'], 'My league')
        self.assertEqual(x['data']['weights']['skater']['goals'], -4)
        self.assertIsNone(x['data']['players'][1]['games'])
        self.assertEqual(x['data']['publication']['status'], 'draft')
        self.assertNotEqual(x['scoringHash'], payload(d,e,d['revision'])['scoringHash'])

    def test_bad_weights_and_revision_fail_before_export(self):
        d, e = fixtures()
        with self.assertRaises(ValueError):payload(d,e,'incorrect')
        for value in ('5', True, float('nan')):
            weights = deepcopy(e['weights']);weights['skater']['goals'] = value
            with self.subTest(value=value), self.assertRaises(ValueError):payload(d,e,d['revision'],weights=weights)

    def test_exported_formulas_cache_unknown_as_blank_and_count_once(self):
        d, e = fixtures()
        weights = deepcopy(e['weights']);weights['skater']['goals'] = 3
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder);source=root/'source.json';editorial=root/'editorial.json';out=root/'review.xlsx'
            source.write_text(json.dumps(d));editorial.write_text(json.dumps(e))
            export(source,editorial,d['revision'],out,league='Goals weighted',weights=weights)
            values=openpyxl.load_workbook(out,data_only=True)
            formulas=openpyxl.load_workbook(out,data_only=False)
            self.assertEqual(values['Players']['G2'].value,10)
            self.assertIsNone(values['Players']['G3'].value)
            self.assertAlmostEqual(values['Players']['AR2'].value,20)
            self.assertIn(values['Players']['AR3'].value,(None,''))
            self.assertAlmostEqual(values['Players']['AF2'].value,2)
            self.assertIn(values['Players']['AF3'].value,(None,''))
            self.assertTrue(formulas['Players']['AR2'].value.startswith('=IF('))
            self.assertNotIn('H2',formulas['Players']['AR2'].value)
            self.assertEqual(values['Read Me']['B2'].value,d['revision'])
            self.assertEqual(values['Scoring']['B5'].value,'Goals weighted')
            self.assertIn('Reviewed team note', [row[1] for row in values['ANA'].iter_rows(values_only=True)])
            with self.assertRaises(ValueError):export(source,editorial,d['revision'],out)


if __name__=='__main__':unittest.main()
